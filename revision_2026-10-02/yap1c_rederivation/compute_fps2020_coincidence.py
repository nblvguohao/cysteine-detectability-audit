"""Re-derivation of the YAP1C (Arabidopsis sulfenylation) site/detection coincidence share.

Supplemental Note 7, paragraph "Re-derivation check for the YAP1C cohort" (post hoc, 2 October 2026).

Computation identical to the run of 2 October 2026 (see coincidence_output.txt); only the input
locations are now configurable instead of hard-coded, and the Wilson helper is inlined.

Inputs (not redistributed; see README.md for sources and SHA-256):
  Ath_TAIR10_pep.fa.gz                 Ensembl Plants TAIR10 pep.all
  Data_1.XLSX, Data_2.xlsx             supplementary data sheets of the YAP1C study (Front. Plant Sci. 2020)
  pLink_reports_{untreated,H2O2}_{proteinLevel,peptideLevel}.zip   PRIDE PXD016723
    (either the four zips or their extracted folders of the same names)

Usage:
  python compute_fps2020_coincidence.py [--input-dir DIR] [--fasta F] [--si-dir DIR] [--plink-dir DIR]
  Defaults: --input-dir is $YAP1C_INPUT_DIR, else ./inputs next to this script; the other three
  default to --input-dir.
Requires: openpyxl.
"""
import argparse
import collections
import csv
import glob
import gzip
import math
import os
import tempfile
import zipfile

import openpyxl


def wilson(k, n, z=1.959964):
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return p, c - h, c + h


ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
ap.add_argument('--input-dir', default=os.environ.get(
    'YAP1C_INPUT_DIR', os.path.join(os.path.dirname(os.path.abspath(__file__)), 'inputs')))
ap.add_argument('--fasta')
ap.add_argument('--si-dir')
ap.add_argument('--plink-dir')
args = ap.parse_args()
FASTA = args.fasta or os.path.join(args.input_dir, 'Ath_TAIR10_pep.fa.gz')
SI = args.si_dir or args.input_dir
PL = args.plink_dir or args.input_dir

# use extracted pLink folders if present, otherwise extract the zips into a temporary folder
PLINK_SETS = ['pLink_reports_%s_%s' % (t, l) for t in ('untreated', 'H2O2')
              for l in ('proteinLevel', 'peptideLevel')]
if not all(os.path.isdir(os.path.join(PL, s)) for s in PLINK_SETS):
    tmp = tempfile.mkdtemp(prefix='yap1c_plink_')
    for s in PLINK_SETS:
        with zipfile.ZipFile(os.path.join(PL, s + '.zip')) as z:
            z.extractall(tmp)
    PL = tmp

seqs = {}
cur = None
for l in gzip.open(FASTA, 'rt'):
    if l[0] == '>':
        cur = l[1:].split()[0]
        seqs[cur] = []
    else:
        seqs[cur].append(l.strip())
seqs = {k: ''.join(v) for k, v in seqs.items()}

# positives: Dataset S2
ws = openpyxl.load_workbook(os.path.join(SI, 'Data_2.xlsx'), read_only=True)['DatasetS2']
pos_all = set(); pos_first = set(); npos = 0
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i < 4 or not r[0] or not r[13]:
        continue
    npos += 1
    prots = str(r[13]).split(';'); sites = str(r[14]).split(';')
    for j, (p, s) in enumerate(zip(prots, sites)):
        pos_all.add((p.strip(), int(s)))
        if j == 0:
            pos_first.add((p.strip(), int(s)))
print('S2 rows', npos, 'pos units all', len(pos_all), 'first', len(pos_first))


def regular(pattern):
    det_all = set(); det_first = set(); npep = 0; ncys = 0; unm = 0; runs = collections.Counter()
    for f in glob.glob(os.path.join(PL, pattern, '*filtered_regular_peptides.csv')):
        for row in csv.reader(open(f)):
            if not row or row[0] == '' or row[0] == 'Peptide_Order':
                if len(row) > 2 and row[0] == '' and row[1].isdigit():
                    runs[row[2].split('.')[0]] += 1
                continue
            pep = row[1]; npep += 1
            if 'C' not in pep:
                continue
            ncys += 1
            prots = [p.strip() for p in row[4].split('/') if p.strip()]
            for j, p in enumerate(prots):
                s = seqs.get(p)
                if s is None:
                    cands = [k for k in seqs if k.split('.')[0] == p.split('.')[0]]
                    s = None
                    for k in cands:
                        if pep in seqs[k]:
                            s = seqs[k]; break
                if s is None or pep not in s:
                    unm += 1; continue
                st = s.find(pep)
                for k, a in enumerate(pep):
                    if a == 'C':
                        det_all.add((p, st + k + 1))
                        if j == 0:
                            det_first.add((p, st + k + 1))
    return det_all, det_first, npep, ncys, unm, runs


for lab, pat in [('A_peptideLevel_same_runs', 'pLink_reports_*_peptideLevel'),
                 ('B_proteinLevel_S_runs', 'pLink_reports_*_proteinLevel'), ('S1b_test_samples', None)]:
    if pat is None:
        continue
    da, dfst, npep, ncys, unm, runs = regular(pat)
    print('==', lab, 'regular peptides', npep, 'Cys peptides', ncys, 'unmapped protein-matches', unm,
          'runs', dict(runs))
    for rule, P, Dt in [('all_matching', pos_all, da), ('first_listed', pos_first, dfst)]:
        U = P | Dt; k = len(P); n = len(U)
        p, lo, hi = wilson(k, n)
        print(f'  {rule}: detected_cys(regular)={len(Dt)} overlap_with_pos={len(P & Dt)} union={n} pos={k} '
              f'coincidence={p:.4f} [{lo:.4f},{hi:.4f}]')

# C) S1b: purification test samples; Bo_1 = nonenriched proteome shotgun
ws = openpyxl.load_workbook(os.path.join(SI, 'Data_1.XLSX'), read_only=True)['DatasetS1b_regular_peptide']
pep2runs = collections.defaultdict(set); pep2prot = {}
curpep = None
for r in ws.iter_rows(min_row=3, values_only=True):
    if isinstance(r[0], int) and isinstance(r[1], str):
        curpep = r[1]; pep2prot[curpep] = [p.strip() for p in str(r[4]).split('/') if p.strip()]
    elif r[0] is None and isinstance(r[1], int) and curpep:
        pep2runs[curpep].add(str(r[2]).split('.')[0])
runsC = collections.Counter(x for v in pep2runs.values() for x in v)
print('S1b runs', runsC)
for run in ['QE_Plus_YangJing_TCP_Bo_1_33per_20171106', 'QE_Plus_YangJing_TCP_Bo_3_33per_20171106']:
    da = set(); dfst = set(); ncys = 0
    for pep, runs in pep2runs.items():
        if run not in runs or 'C' not in pep:
            continue
        ncys += 1
        for j, p in enumerate(pep2prot[pep]):
            s = seqs.get(p)
            if not s or pep not in s:
                continue
            st = s.find(pep)
            for k, a in enumerate(pep):
                if a == 'C':
                    da.add((p, st + k + 1))
                    if j == 0:
                        dfst.add((p, st + k + 1))
    for rule, P, Dt in [('all_matching', pos_all, da), ('first_listed', pos_first, dfst)]:
        ov = len(P & Dt); p, lo, hi = wilson(ov, len(Dt))
        print(f'== C {run} Cys peptides={ncys} {rule}: detected_cys={len(Dt)} positives_among_detected={ov} '
              f'share={p:.4f} [{lo:.4f},{hi:.4f}]')

# independent-readout share for A and B
for lab, pat in [('A', 'pLink_reports_*_peptideLevel'), ('B', 'pLink_reports_*_proteinLevel')]:
    da, dfst, *_ = regular(pat)
    for rule, P, Dt in [('all_matching', pos_all, da), ('first_listed', pos_first, dfst)]:
        ov = len(P & Dt); p, lo, hi = wilson(ov, len(Dt))
        print(f'== {lab} {rule}: positives among regular-detected Cys {ov}/{len(Dt)} = {p:.4f} [{lo:.4f},{hi:.4f}]')
